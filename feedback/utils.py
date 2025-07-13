from pptx import Presentation
from docx import Document
import os
from PIL import Image
import io
import base64

def extract_ppt_content(ppt_file):
    """
    Extract text and image information from a PowerPoint file
    """
    prs = Presentation(ppt_file)
    content = []
    
    for slide_number, slide in enumerate(prs.slides, 1):
        slide_content = {
            'slide_number': slide_number,
            'text': [],
            'images': []
        }
        
      
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_content['text'].append(shape.text.strip())
           
            if shape.shape_type == 13:  
                try:
                    image = shape.image
                    image_bytes = image.blob
                    
                 
                    image_base64 = base64.b64encode(image_bytes).decode('utf-8')
                    slide_content['images'].append({
                        'data': image_base64,
                        'format': image.ext
                    })
                except Exception as e:
                    print(f"Error processing image: {str(e)}")
        
        content.append(slide_content)
    
    return content

def extract_docx_content(docx_file):
    """
    Extract text and image information from a Word document
    """
    doc = Document(docx_file)
    content = []
    
    for para in doc.paragraphs:
        if para.text.strip():
            content.append({
                'type': 'text',
                'content': para.text.strip()
            })
    
  
    for rel in doc.part.rels.values():
        if "image" in rel.target_ref:
            try:
                image_data = rel.target_part.blob
                image_base64 = base64.b64encode(image_data).decode('utf-8')
                content.append({
                    'type': 'image',
                    'data': image_base64,
                    'format': os.path.splitext(rel.target_ref)[1]
                })
            except Exception as e:
                print(f"Error processing image: {str(e)}")
    
    return content

def extract_text_content(text_file):
    """
    Extract content from a text file
    """
    with open(text_file, 'r', encoding='utf-8') as file:
        content = file.read()
    return [{
        'type': 'text',
        'content': content
    }]

def format_content_for_llm(content, file_type):
    """
    Format the extracted content for LLM processing
    """
    if file_type == 'ppt':
        formatted_content = []
        for slide in content:
            slide_text = f"\nSlide {slide['slide_number']}:\n"
            
            if slide['text']:
                slide_text += "Text Content:\n"
                for text in slide['text']:
                    slide_text += f"- {text}\n"
            
            if slide['images']:
                slide_text += f"\nContains {len(slide['images'])} images\n"
            
            formatted_content.append(slide_text)
        return "\n".join(formatted_content)
    
    elif file_type in ['docx', 'txt']:
        formatted_content = []
        for item in content:
            if item['type'] == 'text':
                formatted_content.append(item['content'])
            elif item['type'] == 'image':
                formatted_content.append(f"\n[Image content - {item['format']}]\n")
        return "\n\n".join(formatted_content)
    
    return "" 